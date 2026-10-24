"""1日分（0:00〜23:45）のデータから日次レポート（Excel）を作る．

使い方
  python make_daily_report.py             まだレポートがない「終わった日」をすべて作る（毎日 0:05 ごろ自動実行）
  python make_daily_report.py 20261024    指定した日のレポートを作り直す

出力
  reports/2026/日次レポート_10.24-10.25.xlsx   （10/24 0:00 〜 10/25 0:00 のデータ）

シート
  概要                 主な数値（前日比つき），天気，この日のポイント（文章）
  グラフ               1日の変化のグラフ6つと，それぞれのグラフからわかること
  ステーション別       25箇所の集計表（空車・満車の時間，推定貸出回数など）とコメント
  時間帯別             1時間ごとの集計と，時間帯×ステーションの台数の色分け表
  業者作業の候補       再配置・回収・整備の可能性がある変化の一覧
  ステーション別グラフ 25箇所それぞれの15分ごとの推移
  15分ごとのデータ     この日の全記録
  説明                 用語と判定基準，データの出典
"""
import glob
import os
import sys
from datetime import date, datetime, timedelta, timezone

import pandas as pd
import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from make_excel import (BORDER, CENTER, F_BODY, F_HEAD, F_NOTE, F_TITLE, FILL_EMPTY, FILL_FULL,
                        FILL_HEAD, FILL_MISS, FILL_NAKA, FONT_NAME, LEFT, MED, RIGHT, STATUS_FILL,
                        _txt, build_series, make_line_chart, outline, sheet_charts, style_line,
                        write_chart_data)
from make_summary import WARD

JST = timezone(timedelta(hours=9))
DATA_DIR = "data"
OUT_DIR = "reports"
EXPECTED = 96          # 1日の取得回数（15分ごと）
BIG_CHANGE = 3         # 15分でこの台数以上増減したら「業者作業の候補」
LAT, LON = 35.448, 139.630  # 天気を取る地点（横浜市中区・西区の境付近）
RAW_SHEET = "15分ごとのデータ"
WEEK = "月火水木金土日"

F_H2 = Font(name=FONT_NAME, size=12, bold=True, color="1F4E78")
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
FILL_KEY = PatternFill("solid", fgColor="F2F2F2")
FILL_TEXT = PatternFill("solid", fgColor="FFF9E6")
WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

WMO_JA = {0: "快晴", 1: "晴れ", 2: "晴れ時々曇り", 3: "曇り", 45: "霧", 48: "霧",
          51: "霧雨", 53: "霧雨", 55: "霧雨", 56: "霧雨", 57: "霧雨",
          61: "小雨", 63: "雨", 65: "強い雨", 66: "雨", 67: "雨",
          71: "雪", 73: "雪", 75: "大雪", 77: "雪",
          80: "にわか雨", 81: "にわか雨", 82: "激しいにわか雨", 85: "にわか雪", 86: "にわか雪",
          95: "雷雨", 96: "雷雨", 99: "雷雨"}


# ---------------------------------------------------------------------------
# データの読み込みと分析
# ---------------------------------------------------------------------------
def raw_path(d):
    return f"{DATA_DIR}/yokohama_stations_15min_log_raw_{d:%Y%m%d}.csv"


def load_raw(d):
    p = raw_path(d)
    if not os.path.exists(p):
        return None
    df = pd.read_csv(p, dtype={"station_id": str}, encoding="utf-8-sig")
    df["t"] = pd.to_datetime(df["fetch_time"], format="%Y/%m/%d %H:%M")
    df["missing"] = df["missing"].astype(str).str.lower().eq("true")
    return df.drop_duplicates(subset=["t", "station_id"], keep="last")


def status_of(r):
    if r["missing"]:
        return "欠損"
    if str(r["is_renting"]).lower() == "false":
        return "休止"
    if r["num_docks_available"] == 0:
        return "満車"
    if r["num_bikes_available"] == 0:
        return "空車"
    return "通常"


def fmt_dur(n):
    """15分の回数を「○時間○分」にする．"""
    m = int(n) * 15
    if m == 0:
        return "0分"
    h, mm = divmod(m, 60)
    return (f"{h}時間" if h else "") + (f"{mm}分" if mm else "")


def holiday_name(d):
    try:
        import holidays
        return holidays.country_holidays("JP").get(d)
    except Exception:
        return None


def day_label(d):
    h = holiday_name(d)
    w = WEEK[d.weekday()]
    kind = "祝日" if h else ("日曜" if d.weekday() == 6 else "土曜" if d.weekday() == 5 else "平日")
    return w, kind, h


def fetch_weather(d):
    """Open-Meteo から1時間ごとの気温・降水量・天気を取る．取れなければ None．"""
    params = {"latitude": LAT, "longitude": LON, "timezone": "Asia/Tokyo",
              "hourly": "temperature_2m,precipitation,weather_code",
              "start_date": d.isoformat(), "end_date": d.isoformat()}
    for url in ("https://api.open-meteo.com/v1/forecast",
                "https://archive-api.open-meteo.com/v1/archive"):
        try:
            r = requests.get(url, params=params, timeout=30)
            r.raise_for_status()
            h = r.json()["hourly"]
            w = pd.DataFrame({
                "時": [int(t[11:13]) for t in h["time"]],
                "気温(℃)": h["temperature_2m"],
                "降水量(mm)": h["precipitation"],
                "天気": [WMO_JA.get(c, "") if c is not None else None for c in h["weather_code"]],
            })
            w["気温(℃)"] = w["気温(℃)"].round(1)
            w["降水量(mm)"] = w["降水量(mm)"].round(1)
            if len(w) == 24 and w["気温(℃)"].notna().all():
                return w
        except Exception as e:
            print(f"  天気の取得に失敗（{url}）: {e}")
    return None


def analyze(d):
    """d の日（0:00〜23:59）を分析した結果を辞書で返す．データがなければ None．"""
    cur = load_raw(d)
    if cur is None or cur.empty:
        return None
    prev = load_raw(d - timedelta(days=1))
    if prev is not None and not prev.empty:
        prev = prev[prev["t"] == prev["t"].max()]  # 前日の最後の回（増減の計算にだけ使う）
        df = pd.concat([prev, cur], ignore_index=True)
    else:
        df = cur.copy()

    df["status"] = df.apply(status_of, axis=1)
    df["ward"] = df["station_id"].map(WARD).fillna("")
    names = df.dropna(subset=["name"]).groupby("station_id")["name"].last()
    caps = df.dropna(subset=["capacity"]).groupby("station_id")["capacity"].last()
    df["name"] = df["name"].fillna(df["station_id"].map(names))
    df["capacity"] = df["capacity"].fillna(df["station_id"].map(caps))
    df["occ"] = df["capacity"] - df["num_docks_available"]  # 枠に入っている台数

    # 前の回からの増減（欠損の回は除き，間隔が20分を超えたら数えない）
    ok = df[df["status"] != "欠損"].sort_values(["station_id", "t"]).copy()
    g = ok.groupby("station_id")
    ok["dt"] = g["t"].diff().dt.total_seconds() / 60
    for c in ["occ", "num_bikes_available", "unavailable"]:
        ok["d_" + c] = g[c].diff().where(ok["dt"] <= 20)
    df = df.merge(ok[["t", "station_id", "d_occ", "d_num_bikes_available", "d_unavailable"]],
                  on=["t", "station_id"], how="left")
    df = df[df["t"].dt.date == d].sort_values(["t", "station_id"]).reset_index(drop=True)
    df["hour"] = df["t"].dt.hour
    ok = df[df["status"] != "欠損"]

    info = (df.groupby("station_id")
              .agg(区=("ward", "last"), ステーション名=("name", "last"), 容量=("capacity", "last"))
              .reset_index().rename(columns={"station_id": "ステーションID"})
              .sort_values(["区", "ステーションID"]).reset_index(drop=True))

    # ステーション別
    st = ok.groupby("station_id").agg(
        n=("status", "size"),
        empty=("status", lambda s: (s == "空車").sum()),
        full=("status", lambda s: (s == "満車").sum()),
        avail_mean=("num_bikes_available", "mean"),
        avail_min=("num_bikes_available", "min"),
        avail_max=("num_bikes_available", "max"),
        unav_mean=("unavailable", "mean"),
        unav_min=("unavailable", "min"),
        rent=("d_occ", lambda s: -s[s < 0].sum()),
        ret=("d_occ", lambda s: s[s > 0].sum()),
    )
    st = info.set_index("ステーションID").join(st).reset_index()
    st[["n", "empty", "full", "rent", "ret"]] = st[["n", "empty", "full", "rent", "ret"]].fillna(0)

    # 回ごと（25箇所の合計）
    per = df.groupby("t").agg(
        empty=("status", lambda s: (s == "空車").sum()),
        full=("status", lambda s: (s == "満車").sum()),
        avail=("num_bikes_available", "sum"),
        unav=("unavailable", "sum"),
        miss=("status", lambda s: (s == "欠損").sum()),
    ).reset_index()

    # 時間帯別
    hr = pd.DataFrame({"時": range(24)})
    ph = per.assign(hour=per["t"].dt.hour).groupby("hour")
    hr["取得回数"] = hr["時"].map(ph.size()).fillna(0).astype(int)
    hr["空車ステーション数(平均)"] = hr["時"].map(ph["empty"].mean()).round(1)
    hr["満車ステーション数(平均)"] = hr["時"].map(ph["full"].mean()).round(1)
    hr["合計貸出可能台数(平均)"] = hr["時"].map(ph["avail"].mean()).round(1)
    hr["合計貸出不可台数(平均)"] = hr["時"].map(ph["unav"].mean()).round(1)
    oh = ok.groupby("hour")["d_occ"]
    hr["推定貸出回数"] = hr["時"].map(oh.apply(lambda s: -s[s < 0].sum())).fillna(0).astype(int)
    hr["推定返却回数"] = hr["時"].map(oh.apply(lambda s: s[s > 0].sum())).fillna(0).astype(int)

    heat = ok.pivot_table(index="hour", columns="station_id", values="num_bikes_available", aggfunc="mean")
    heat = heat.reindex(index=range(24), columns=info["ステーションID"]).round(1)

    # 業者作業の候補
    ev = []
    for _, r in ok.iterrows():
        do, da, du = r["d_occ"], r["d_num_bikes_available"], r["d_unavailable"]
        if pd.isna(do):
            continue
        if abs(do) >= BIG_CHANGE:
            if do > 0:
                kind, note = "大きな増加", "業者による搬入（再配置）の可能性．利用者のまとまった返却の場合もある"
            elif du <= -BIG_CHANGE and da >= 0:
                kind, note = "大きな減少", "貸出不可の車両だけがまとめて減った．故障車などの回収の可能性"
            else:
                kind, note = "大きな減少", "業者による持ち出し（再配置）の可能性．利用者のまとまった貸出の場合もある"
        elif do == 0 and du <= -1 and da >= 1:
            kind, note = "整備らしい動き", "台数は同じまま貸出不可→貸出可能．バッテリー交換・車両の入れ替え・予約の取り消しのいずれか"
        else:
            continue
        ev.append({"日時": r["t"].strftime("%H:%M"), "区": r["ward"], "ステーションID": r["station_id"],
                   "ステーション名": r["name"], "種類": kind,
                   "枠内の増減": int(do), "貸出可能の増減": int(da), "貸出不可の増減": int(du), "解釈": note})
    events = pd.DataFrame(ev, columns=["日時", "区", "ステーションID", "ステーション名", "種類",
                                       "枠内の増減", "貸出可能の増減", "貸出不可の増減", "解釈"])

    return {"date": d, "df": df, "ok": ok, "info": info, "st": st, "per": per, "hr": hr,
            "heat": heat, "events": events, "n_fetch": df["t"].nunique(),
            "n_missing": int((df["status"] == "欠損").sum())}


def totals(a):
    st = a["st"]
    return {
        "取得回数": a["n_fetch"],
        "欠損（件）": a["n_missing"],
        "空車の延べ時間": int(st["empty"].sum()),
        "満車の延べ時間": int(st["full"].sum()),
        "推定貸出回数": int(st["rent"].sum()),
        "推定返却回数": int(st["ret"].sum()),
        "貸出不可台数（25箇所合計の平均）": round(float(a["per"]["unav"].mean()), 1),
        "業者作業の候補（件）": len(a["events"]),
    }


# ---------------------------------------------------------------------------
# 文章（グラフからわかること）
# ---------------------------------------------------------------------------
def hour_label(h):
    return f"{int(h)}時台"


def texts(a, w, prev_tot):
    st, hr, per, ev = a["st"], a["hr"], a["per"], a["events"]
    nm = lambda r: r["ステーション名"]
    out = {}

    # ① 空車・満車ステーション数
    e = hr[hr["取得回数"] > 0]
    emax, emin = e.loc[e["空車ステーション数(平均)"].idxmax()], e.loc[e["空車ステーション数(平均)"].idxmin()]
    pmax = per.loc[per["empty"].idxmax()]
    t1 = [f"空車のステーションは1日平均 {per['empty'].mean():.1f} 箇所（25箇所中）でした．",
          f"最も多かったのは {hour_label(emax['時'])}（平均 {emax['空車ステーション数(平均)']} 箇所），"
          f"最も少なかったのは {hour_label(emin['時'])}（平均 {emin['空車ステーション数(平均)']} 箇所）です．"
          f"1回の記録での最多は {pmax['t']:%H:%M} の {int(pmax['empty'])} 箇所です．"]
    if per["full"].max() > 0:
        ft = per[per["full"] > 0]["t"]
        t1.append(f"満車のステーションは {ft.min():%H:%M}〜{ft.max():%H:%M} の間に現れ，"
                  f"最大で同時に {int(per['full'].max())} 箇所でした．")
    else:
        t1.append("満車のステーションはこの日一度もありませんでした．")
    out["empty_full"] = t1

    # ② 合計の貸出可能・貸出不可
    am, an = per.loc[per["avail"].idxmax()], per.loc[per["avail"].idxmin()]
    um, un = per.loc[per["unav"].idxmax()], per.loc[per["unav"].idxmin()]
    t2 = [f"25箇所合計の貸出可能台数は {am['t']:%H:%M} に最大 {int(am['avail'])} 台，"
          f"{an['t']:%H:%M} に最小 {int(an['avail'])} 台でした．",
          f"貸出不可の台数（故障・電池切れ・予約中などの合計）は {um['t']:%H:%M} に最大 {int(um['unav'])} 台，"
          f"{un['t']:%H:%M} に最小 {int(un['unav'])} 台でした．"]
    seg = lambda h0, h1: per[per["t"].dt.hour.between(h0, h1)]["unav"].mean()
    morning, midday, night = seg(5, 8), seg(11, 15), seg(20, 23)
    if pd.notna(morning) and pd.notna(midday) and midday <= morning - 2:
        t2.append(f"貸出不可は早朝（5〜8時，平均 {morning:.1f} 台）より昼（11〜15時，平均 {midday:.1f} 台）の方が少なく，"
                  "日中にバッテリー交換などの整備が行われた可能性があります．")
    if pd.notna(midday) and pd.notna(night) and night >= midday + 2:
        t2.append(f"夜（20〜23時）には平均 {night:.1f} 台に増えており，電池が減った車両の返却などで"
                  "貸せない車両が夜にたまる傾向がみられます．")
    out["totals"] = t2

    # ③ 時間帯別の貸出・返却
    rent, ret = int(hr["推定貸出回数"].sum()), int(hr["推定返却回数"].sum())
    rm = hr.loc[hr["推定貸出回数"].idxmax()]
    rr = hr.loc[hr["推定返却回数"].idxmax()]
    t3 = [f"推定貸出回数は1日で {rent} 回，推定返却回数は {ret} 回でした（15分ごとの台数の減少・増加から数えた下限の値）．",
          f"貸出が最も多かったのは {hour_label(rm['時'])}（{int(rm['推定貸出回数'])} 回），"
          f"返却が最も多かったのは {hour_label(rr['時'])}（{int(rr['推定返却回数'])} 回）です．"]
    am_r = hr[hr["時"].between(7, 9)]["推定貸出回数"].sum()
    pm_r = hr[hr["時"].between(17, 19)]["推定貸出回数"].sum()
    night_r = hr[hr["時"].between(0, 4)]["推定貸出回数"].sum()
    if rent > 0:
        t3.append(f"朝（7〜9時）の貸出は {am_r} 回，夕方（17〜19時）は {pm_r} 回，深夜（0〜4時）は {night_r} 回でした．")
    out["hourly"] = t3

    # ④ 天気
    if w is not None:
        rain = w["降水量(mm)"].fillna(0)
        rh = w[rain >= 0.5]["時"].tolist()
        t4 = [f"気温は最高 {w['気温(℃)'].max():.1f}℃，最低 {w['気温(℃)'].min():.1f}℃，"
              f"降水量は1日合計 {rain.sum():.1f} mm でした．"]
        if rh:
            t4.append(f"0.5mm以上の雨が降ったのは {len(rh)} 時間（{', '.join(str(h) + '時' for h in rh[:8])}"
                      f"{' ほか' if len(rh) > 8 else ''}）です．")
            m = hr.merge(w, on="時")
            wet = m[m["降水量(mm)"] >= 0.5]["推定貸出回数"].mean()
            dry = m[(m["降水量(mm)"] < 0.5) & m["時"].between(6, 22)]["推定貸出回数"].mean()
            if pd.notna(wet) and pd.notna(dry):
                t4.append(f"雨の時間帯の推定貸出回数は1時間平均 {wet:.1f} 回，雨のない6〜22時は {dry:.1f} 回でした．")
        else:
            t4.append("1日を通して雨はほとんど降りませんでした．")
        out["weather"] = t4

    # ⑤ ステーション別の空車・満車時間
    s = st.sort_values("empty", ascending=False)
    top = s.head(3)
    t5 = ["空車の時間が長かったのは " + "，".join(f"{nm(r)}（{fmt_dur(r['empty'])}）" for _, r in top.iterrows()) + " です．"]
    allday = st[(st["n"] > 0) & (st["empty"] >= st["n"] * 0.95)]
    if len(allday):
        t5.append("ほぼ終日空車だったのは " + "，".join(allday["ステーション名"]) + " で，補充（再配置）の必要性が高いステーションです．")
    noempty = st[st["empty"] == 0]
    if len(noempty):
        t5.append("一度も空車にならなかったのは " + "，".join(noempty["ステーション名"]) + " です．")
    full = st[st["full"] > 0].sort_values("full", ascending=False)
    if len(full):
        t5.append("満車になったのは " + "，".join(f"{nm(r)}（{fmt_dur(r['full'])}）" for _, r in full.iterrows())
                  + " で，この時間は返却したい利用者が返却できない状態でした．")
    else:
        t5.append("満車になったステーションはありませんでした．")
    out["station_ef"] = t5

    # ⑥ ステーション別の貸出・返却
    s = st.sort_values("rent", ascending=False)
    t6 = ["推定貸出回数が多かったのは " + "，".join(f"{nm(r)}（{int(r['rent'])}回）" for _, r in s.head(3).iterrows()) + " です．"]
    zero = st[st["rent"] == 0]
    if len(zero):
        t6.append(f"貸出が一度も記録されなかったステーションは {len(zero)} 箇所（" + "，".join(zero["ステーション名"].head(6))
                  + ("ほか" if len(zero) > 6 else "") + "）で，その多くは自転車が置かれていない（空車の）時間が長いことが原因と考えられます．")
    st2 = st.assign(diff=st["ret"] - st["rent"])
    inflow = st2.sort_values("diff", ascending=False).iloc[0]
    outflow = st2.sort_values("diff").iloc[0]
    if inflow["diff"] > 0:
        t6.append(f"返却が貸出を最も上回ったのは {nm(inflow)}（＋{int(inflow['diff'])}台）で，自転車がたまりやすい場所です．")
    if outflow["diff"] < 0:
        t6.append(f"貸出が返却を最も上回ったのは {nm(outflow)}（{int(outflow['diff'])}台）で，自転車が出ていきやすい場所です．")
    out["station_use"] = t6

    # 概要（この日のポイント）
    tot = totals(a)
    pts = []
    q = "そろっています" if tot["取得回数"] >= EXPECTED and tot["欠損（件）"] == 0 else "一部が欠けています"
    pts.append(f"【データ】取得 {tot['取得回数']} 回（予定 {EXPECTED} 回），欠損 {tot['欠損（件）']} 件で，この日のデータは{q}．")
    pts.append("【空車・満車】" + t1[0] + t1[1] + (t1[2] if len(t1) > 2 else ""))
    pts.append("【利用】" + t3[0] + t3[1])
    pts.append("【空車が長いステーション】" + t5[0] + (t5[1] if len(allday) else ""))
    if len(full):
        pts.append("【満車】" + [x for x in t5 if x.startswith("満車になった")][0])
    stuck = st[st["unav_min"] >= 2].sort_values("unav_min", ascending=False)
    if len(stuck):
        pts.append("【長時間の貸出不可】" + "，".join(f"{nm(r)}（終日{int(r['unav_min'])}台以上）" for _, r in stuck.iterrows())
                   + " は，1日中貸出不可の車両が残っていました．予約では説明しにくく，故障などで長く動いていない車両がある可能性があります．")
    big = ev[ev["種類"] != "整備らしい動き"]
    mnt = ev[ev["種類"] == "整備らしい動き"]
    if len(big):
        pts.append("【業者作業の候補】15分で3台以上の増減が " + f"{len(big)} 件ありました（"
                   + "，".join(f"{r['日時']} {r['ステーション名']} {r['枠内の増減']:+d}台" for _, r in big.head(4).iterrows())
                   + ("ほか" if len(big) > 4 else "") + "）．詳しくは「業者作業の候補」シートを参照．")
    else:
        pts.append("【業者作業の候補】15分で3台以上動いた変化はなく，トラックによる再配置の形跡は見られませんでした．")
    if len(mnt):
        pts.append(f"【整備らしい動き】台数が変わらず貸出不可が貸出可能に変わった変化が {len(mnt)} 件ありました"
                   f"（{', '.join(mnt['日時'].head(6))}{' ほか' if len(mnt) > 6 else ''}）．")
    if w is not None:
        pts.append("【天気】" + " ".join(out["weather"][:2]))
    if prev_tot:
        def chg(k, unit=""):
            cur, pv = tot[k], prev_tot[k]
            return f"{pv}→{cur}{unit}（{cur - pv:+d}{unit}）" if isinstance(cur, int) else f"{pv}→{cur}"
        pts.append(f"【前日との比較】推定貸出回数 {chg('推定貸出回数', '回')}，空車の延べ時間 "
                   f"{fmt_dur(prev_tot['空車の延べ時間'])}→{fmt_dur(tot['空車の延べ時間'])}，満車の延べ時間 "
                   f"{fmt_dur(prev_tot['満車の延べ時間'])}→{fmt_dur(tot['満車の延べ時間'])}．")
    out["points"] = pts
    return out


def station_comment(r):
    c = []
    if r["n"] == 0:
        return "データなし"
    er = r["empty"] / r["n"]
    if er >= 0.95:
        c.append("ほぼ終日空車．補充の優先度が高い")
    elif er >= 0.5:
        c.append(f"空車が長い（{er:.0%}）")
    if r["full"] > 0:
        c.append(f"満車 {fmt_dur(r['full'])}．返却できない時間あり")
    if r["unav_min"] >= 2:
        c.append(f"終日{int(r['unav_min'])}台以上が貸出不可．動いていない車両の可能性")
    if r["rent"] >= 8:
        c.append("利用が多い")
    return "，".join(c) if c else "大きな問題なし"


# ---------------------------------------------------------------------------
# Excel の部品
# ---------------------------------------------------------------------------
def head_cell(c):
    c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER


def body_cell(c, align=RIGHT, fmt=None):
    c.font, c.border, c.alignment = F_BODY, BORDER, align
    if fmt:
        c.number_format = fmt


def sheet_title(ws, text, note=None):
    ws.sheet_view.showGridLines = False
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE


def text_lines(ws, row, col, ncols, lines, width_units, fill=None):
    """文章を1行ずつ結合セルに書き，行の高さを文字数に合わせる．次の行番号を返す．"""
    per_line = max(10, int(width_units / 1.9))  # 1行に入る全角文字数の目安
    for s in lines:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + ncols - 1)
        c = ws.cell(row=row, column=col, value="・" + s)
        c.font, c.alignment = F_BODY, WRAP
        if fill:
            for j in range(col, col + ncols):
                ws.cell(row=row, column=j).fill = fill
        n = -(-len(s) // per_line)
        ws.row_dimensions[row].height = max(16, 15 * n + 3)
        row += 1
    return row


def h2(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_H2
    return row + 1


def fix_layout(ch, x=0.07, y=0.13, w=0.90, h=0.55):
    ch.plot_area.layout = Layout(manualLayout=ManualLayout(
        layoutTarget="inner", xMode="edge", yMode="edge", x=x, y=y, w=w, h=h))


def bar_chart(title_text, horizontal=False, width=24, height=9):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.overlap = -10
    ch.title = title_text
    from openpyxl.drawing.text import CharacterProperties, ParagraphProperties
    for para in ch.title.tx.rich.p:
        para.pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1100, b=True))
        for run in para.r or []:
            run.rPr = CharacterProperties(sz=1100, b=True)
    ch.title.overlay = False
    ch.width, ch.height = width, height
    ch.y_axis.majorGridlines = ChartLines()
    ch.y_axis.scaling.min = 0
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.y_axis.txPr = _txt(800)
    ch.x_axis.txPr = _txt(800)
    ch.legend.position = "b"
    ch.legend.overlay = False
    ch.legend.txPr = _txt(900)
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"  # 上から表の順に並べる
        ch.x_axis.tickLblSkip = 1
        ch.y_axis.crosses = "max"                 # 数値の目盛りを下側に置く
        fix_layout(ch, x=0.30, y=0.08, w=0.66, h=0.78)
    else:
        fix_layout(ch, x=0.07, y=0.13, w=0.90, h=0.62)
    return ch


# ---------------------------------------------------------------------------
# シート
# ---------------------------------------------------------------------------
def sheet_overview(wb, a, w, tx, tot, prev_tot):
    ws = wb.active
    ws.title = "概要"
    d = a["date"]
    wday, kind, hname = day_label(d)
    e = d + timedelta(days=1)
    sheet_title(ws, f"日次レポート：{d.year}年{d.month}月{d.day}日（{wday}）{' ' + hname if hname else ''}",
                f"集計期間：{d.month}/{d.day} 0:00 〜 {e.month}/{e.day} 0:00（15分ごと）　対象：横浜市中区・西区のHELLO CYCLING 25ステーション　"
                f"区分：{kind}　自動作成：{datetime.now(JST):%Y/%m/%d %H:%M}")
    for col, wdt in zip("ABCDEFGHIJKL", [30, 13, 13, 13, 4, 24, 13, 13, 13, 13, 13, 13]):
        ws.column_dimensions[col].width = wdt

    # 主な数値
    r = h2(ws, 4, 1, "主な数値")
    for j, t in enumerate(["項目", "この日", "前日", "前日との差"], 1):
        head_cell(ws.cell(row=r, column=j, value=t))
    keys = list(tot.keys())
    for i, k in enumerate(keys, r + 1):
        v, pv = tot[k], (prev_tot or {}).get(k)
        dur = k.endswith("延べ時間")
        body_cell(ws.cell(row=i, column=1, value=k), LEFT)
        ws.cell(row=i, column=1).fill = FILL_KEY
        body_cell(ws.cell(row=i, column=2, value=fmt_dur(v) if dur else v))
        body_cell(ws.cell(row=i, column=3, value=(fmt_dur(pv) if dur else pv) if pv is not None else "－"))
        if pv is not None:
            diff = v - pv
            txt = (("＋" if diff >= 0 else "－") + fmt_dur(abs(diff))) if dur else (round(diff, 1) if isinstance(diff, float) else diff)
            c = ws.cell(row=i, column=4, value=txt)
            body_cell(c, fmt=None if dur else '+0.0;-0.0;0' if isinstance(diff, float) else '+0;-0;0')
        else:
            body_cell(ws.cell(row=i, column=4, value="－"))
    end = r + len(keys)
    outline(ws, r, 1, end, 4)
    ws.cell(row=end + 1, column=1, value="延べ時間＝25箇所それぞれの時間の合計．推定回数は15分ごとの増減から数えた下限の値．").font = F_NOTE

    # 天気
    r2 = h2(ws, 4, 6, "天気（横浜市中区・西区付近）")
    if w is not None:
        rain = w["降水量(mm)"].fillna(0)
        day = w[w["時"].between(7, 18)]["天気"].dropna()
        rows = [("日中（7〜18時）の天気", day.mode().iloc[0] if len(day) else "－"),
                ("最高気温", f"{w['気温(℃)'].max():.1f} ℃"),
                ("最低気温", f"{w['気温(℃)'].min():.1f} ℃"),
                ("降水量（1日合計）", f"{rain.sum():.1f} mm"),
                ("雨の時間（0.5mm以上）", f"{int((rain >= 0.5).sum())} 時間"),
                ("出典", "Open-Meteo.com（CC BY 4.0）")]
    else:
        rows = [("天気", "取得できませんでした")]
    for i, (k, v) in enumerate(rows, r2):
        ws.merge_cells(start_row=i, start_column=7, end_row=i, end_column=8)
        body_cell(ws.cell(row=i, column=6, value=k), LEFT)
        ws.cell(row=i, column=6).fill = FILL_KEY
        body_cell(ws.cell(row=i, column=7, value=v), LEFT)
        body_cell(ws.cell(row=i, column=8))
    outline(ws, r2, 6, r2 + len(rows) - 1, 8)

    # この日のポイント
    r = max(end, r2 + len(rows)) + 3
    r = h2(ws, r, 1, "この日のポイント")
    width = sum(ws.column_dimensions[c].width for c in "ABCDEFGHIJKL")
    r = text_lines(ws, r, 1, 12, tx["points"], width, fill=FILL_TEXT)
    ws.cell(row=r + 1, column=1, value="※文章はデータから自動で作成しています．グラフは「グラフ」シート，判定の基準は「説明」シートを参照してください．").font = F_NOTE


def sheet_graphs(wb, a, w, tx, gd):
    """gd：グラフ用データのシートと各ブロックの位置．"""
    ws = wb.create_sheet("グラフ")
    sheet_title(ws, "1日の変化のグラフ", "左がグラフ，右がそのグラフからわかること（データから自動で作成）．")
    for j in range(1, 15):
        ws.column_dimensions[get_column_letter(j)].width = 9.5
    ws.column_dimensions["O"].width = 2
    for j in range(16, 24):
        ws.column_dimensions[get_column_letter(j)].width = 11
    text_width = 11 * 8
    ds = gd["ws"]
    row = 4

    def place(ch, heading, lines, rows_needed=21):
        nonlocal row
        h2(ws, row, 1, heading)
        ws.add_chart(ch, f"A{row + 2}")
        c = ws.cell(row=row, column=16, value="このグラフからわかること")
        c.font = F_BOLD
        end = text_lines(ws, row + 1, 16, 8, lines, text_width, fill=FILL_TEXT)
        row = max(row + rows_needed, end + 2)

    n = gd["n_per"]
    cats = Reference(ds, min_col=1, min_row=2, max_row=n + 1)
    skip = max(1, -(-n // 24))

    ch = make_line_chart("空車・満車になっているステーション数（15分ごと）", skip, width=24)
    for col, color in ((2, "2E75B6"), (3, "C00000")):
        ch.add_data(Reference(ds, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        style_line(ch.series[-1], color, width=25400)
    ch.set_categories(cats)
    place(ch, "① 空車・満車のステーション数", tx["empty_full"])

    ch = make_line_chart("25箇所合計の貸出可能台数・貸出不可台数（15分ごと）", skip, width=24)
    for col, color in ((4, "548235"), (5, "C00000")):
        ch.add_data(Reference(ds, min_col=col, min_row=1, max_row=n + 1), titles_from_data=True)
        style_line(ch.series[-1], color, width=25400)
    ch.set_categories(cats)
    place(ch, "② 貸せる自転車と貸せない自転車の台数", tx["totals"])

    hc = gd["hour_col"]
    ch = bar_chart("時間帯別の推定貸出回数・返却回数")
    for off, color in ((1, "ED7D31"), (2, "4472C4")):
        ch.add_data(Reference(ds, min_col=hc + off, min_row=1, max_row=25), titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
        ch.series[-1].graphicalProperties.line.solidFill = color
    ch.set_categories(Reference(ds, min_col=hc, min_row=2, max_row=25))
    place(ch, "③ 時間帯別の利用（推定）", tx["hourly"])

    if w is not None:
        bar = bar_chart("時間帯別の降水量（棒）と気温（線）")
        bar.add_data(Reference(ds, min_col=hc + 4, min_row=1, max_row=25), titles_from_data=True)
        bar.series[0].graphicalProperties.solidFill = "5B9BD5"
        bar.y_axis.title = None
        bar.set_categories(Reference(ds, min_col=hc, min_row=2, max_row=25))
        line = LineChart()
        line.add_data(Reference(ds, min_col=hc + 3, min_row=1, max_row=25), titles_from_data=True)
        style_line(line.series[0], "C00000", width=25400)
        line.y_axis.axId = 200
        line.y_axis.crosses = "max"
        line.y_axis.delete = False
        line.y_axis.majorGridlines = None
        line.y_axis.txPr = _txt(800)
        bar += line
        fix_layout(bar, x=0.07, y=0.13, w=0.86, h=0.62)
        place(bar, "④ 天気（降水量は左の目盛り mm，気温は右の目盛り ℃）", tx["weather"])

    sc, ns = gd["st_col"], gd["n_st"]
    ch = bar_chart("ステーション別の空車時間・満車時間（時間）", horizontal=True, height=16)
    for off, color in ((1, "5B9BD5"), (2, "E46C0A")):
        ch.add_data(Reference(ds, min_col=sc + off, min_row=1, max_row=ns + 1), titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
    ch.set_categories(Reference(ds, min_col=sc, min_row=2, max_row=ns + 1))
    place(ch, "⑤ ステーション別の空車・満車の時間", tx["station_ef"], rows_needed=35)

    ch = bar_chart("ステーション別の推定貸出回数・返却回数", horizontal=True, height=16)
    for off, color in ((3, "ED7D31"), (4, "4472C4")):
        ch.add_data(Reference(ds, min_col=sc + off, min_row=1, max_row=ns + 1), titles_from_data=True)
        ch.series[-1].graphicalProperties.solidFill = color
    ch.set_categories(Reference(ds, min_col=sc, min_row=2, max_row=ns + 1))
    place(ch, "⑥ ステーション別の利用（推定）", tx["station_use"], rows_needed=35)


def sheet_graph_data(wb, a, w):
    ws = wb.create_sheet("グラフ用データ")
    ws.sheet_properties.tabColor = "A6A6A6"
    per, hr, st = a["per"], a["hr"], a["st"]
    # 回ごと（A〜E）
    for j, t in enumerate(["時刻", "空車ステーション数", "満車ステーション数", "合計貸出可能台数", "合計貸出不可台数"], 1):
        head_cell(ws.cell(row=1, column=j, value=t))
    for i, r in enumerate(per.itertuples(), 2):
        for j, v in enumerate([r.t.strftime("%H:%M"), int(r.empty), int(r.full), int(r.avail), int(r.unav)], 1):
            ws.cell(row=i, column=j, value=v)
    # 時間帯別（G〜K）
    hc = 7
    for j, t in enumerate(["時", "推定貸出回数", "推定返却回数", "気温(℃)", "降水量(mm)"]):
        head_cell(ws.cell(row=1, column=hc + j, value=t))
    wmap = w.set_index("時") if w is not None else None
    for i, r in enumerate(hr.itertuples(), 2):
        ws.cell(row=i, column=hc, value=f"{r.時}時")
        ws.cell(row=i, column=hc + 1, value=int(r.推定貸出回数))
        ws.cell(row=i, column=hc + 2, value=int(r.推定返却回数))
        if wmap is not None:
            ws.cell(row=i, column=hc + 3, value=wmap.loc[r.時, "気温(℃)"])
            ws.cell(row=i, column=hc + 4, value=wmap.loc[r.時, "降水量(mm)"])
    # ステーション別（M〜Q）
    sc = 13
    for j, t in enumerate(["ステーション", "空車時間", "満車時間", "推定貸出回数", "推定返却回数"]):
        head_cell(ws.cell(row=1, column=sc + j, value=t))
    for i, r in enumerate(st.itertuples(), 2):
        ws.cell(row=i, column=sc, value=f"{r.区} {r.ステーション名}")
        ws.cell(row=i, column=sc + 1, value=round(r.empty * 0.25, 2))
        ws.cell(row=i, column=sc + 2, value=round(r.full * 0.25, 2))
        ws.cell(row=i, column=sc + 3, value=int(r.rent))
        ws.cell(row=i, column=sc + 4, value=int(r.ret))
    for j in range(1, sc + 5):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.column_dimensions[get_column_letter(sc)].width = 40
    return {"ws": ws, "n_per": len(per), "hour_col": hc, "st_col": sc, "n_st": len(st)}


def sheet_raw(wb, a):
    ws = wb.create_sheet(RAW_SHEET)
    sheet_title(ws, "15分ごとの全記録", "枠内の台数＝容量−空き枠（貸せる＋貸せない自転車）．増減は前の回（15分前）との差．")
    df = a["df"]
    cols = ["日時", "時刻", "時", "区", "ステーションID", "ステーション名", "容量", "貸出可能台数", "空き枠",
            "貸出不可台数", "枠内の台数", "枠内の増減", "状態"]
    widths = [17, 7, 5, 6, 12, 40, 7, 12, 8, 12, 10, 10, 8]
    for j, (t, wd) in enumerate(zip(cols, widths), 1):
        head_cell(ws.cell(row=4, column=j, value=t))
        ws.column_dimensions[get_column_letter(j)].width = wd
    for i, r in enumerate(df.itertuples(), 5):
        vals = [r.t.strftime("%Y/%m/%d %H:%M"), r.t.strftime("%H:%M"), r.hour, r.ward, r.station_id, r.name,
                r.capacity, r.num_bikes_available, r.num_docks_available, r.unavailable, r.occ, r.d_occ, r.status]
        for j, v in enumerate(vals, 1):
            v = None if (not isinstance(v, str) and pd.isna(v)) else v
            if isinstance(v, float) and float(v).is_integer():
                v = int(v)
            c = ws.cell(row=i, column=j, value=v)
            body_cell(c, LEFT if j in (1, 2, 4, 5, 6, 13) else RIGHT, fmt="+0;-0;0" if j == 12 else None)
    end = 4 + len(df)
    for name, fill in STATUS_FILL.items():
        ws.conditional_formatting.add(f"A5:M{end}", FormulaRule(formula=[f'$M5="{name}"'], fill=fill))
    ws.conditional_formatting.add(f"L5:L{end}", CellIsRule(operator="greaterThanOrEqual", formula=[str(BIG_CHANGE)],
                                                            font=Font(name=FONT_NAME, size=10, bold=True, color="C00000")))
    ws.conditional_formatting.add(f"L5:L{end}", CellIsRule(operator="lessThanOrEqual", formula=[str(-BIG_CHANGE)],
                                                            font=Font(name=FONT_NAME, size=10, bold=True, color="C00000")))
    outline(ws, 4, 1, end, len(cols))
    ws.auto_filter.ref = f"A4:M{end}"
    ws.freeze_panes = "A5"


def sheet_stations(wb, a):
    ws = wb.create_sheet("ステーション別")
    sheet_title(ws, "ステーション別の集計",
                "数値は「15分ごとのデータ」シートから数式で計算しています．空車率が高いほど青，満車率が高いほど赤が濃くなります．")
    cols = ["区", "ステーションID", "ステーション名", "容量", "取得回数", "空車回数", "空車時間(時間)", "空車率",
            "満車回数", "満車時間(時間)", "満車率", "平均貸出可能台数", "最小貸出可能台数", "最大貸出可能台数",
            "平均貸出不可台数", "終日続いた貸出不可台数", "推定貸出回数", "推定返却回数", "回転率(回/台)", "コメント"]
    widths = [6, 12, 44, 7, 9, 9, 10, 8, 9, 10, 8, 10, 10, 10, 10, 12, 10, 10, 10, 75]
    for j, (t, wd) in enumerate(zip(cols, widths), 1):
        head_cell(ws.cell(row=4, column=j, value=t))
        ws.column_dimensions[get_column_letter(j)].width = wd
    ws.row_dimensions[4].height = 42
    R = f"'{RAW_SHEET}'!"
    ID, ST, AV, UN, DO = R + "$E:$E", R + "$M:$M", R + "$H:$H", R + "$J:$J", R + "$L:$L"
    st = a["st"]
    for i, r in enumerate(st.itertuples(), 5):
        b = f"$B{i}"
        f = [r.区, r.ステーションID, r.ステーション名, r.容量,
             f'=COUNTIFS({ID},{b},{ST},"<>欠損")',
             f'=COUNTIFS({ID},{b},{ST},"空車")', f"=F{i}*0.25", f'=IF(E{i}=0,"",F{i}/E{i})',
             f'=COUNTIFS({ID},{b},{ST},"満車")', f"=I{i}*0.25", f'=IF(E{i}=0,"",I{i}/E{i})',
             f'=IFERROR(AVERAGEIFS({AV},{ID},{b},{ST},"<>欠損"),"")',
             f'=IF(E{i}=0,"",_xlfn.MINIFS({AV},{ID},{b},{ST},"<>欠損"))',
             f'=IF(E{i}=0,"",_xlfn.MAXIFS({AV},{ID},{b},{ST},"<>欠損"))',
             f'=IFERROR(AVERAGEIFS({UN},{ID},{b},{ST},"<>欠損"),"")',
             f'=IF(E{i}=0,"",_xlfn.MINIFS({UN},{ID},{b},{ST},"<>欠損"))',
             f'=-SUMIFS({DO},{ID},{b},{DO},"<0")', f'=SUMIFS({DO},{ID},{b},{DO},">0")',
             f'=IF(D{i}=0,"",Q{i}/D{i})', station_comment(st.iloc[i - 5])]
        fmts = {7: "0.00", 8: "0%", 10: "0.00", 11: "0%", 12: "0.0", 15: "0.0", 19: "0.00"}
        for j, v in enumerate(f, 1):
            c = ws.cell(row=i, column=j, value=v)
            body_cell(c, WRAP if j == 20 else LEFT if j in (1, 2, 3) else RIGHT, fmts.get(j))
            if r.区 == "中区":
                c.fill = FILL_NAKA
    end = 4 + len(st)
    t = end + 1
    body_cell(ws.cell(row=t, column=3, value="合計（25箇所）"), LEFT)
    for col in "FGIJQR":
        c = ws[f"{col}{t}"]
        c.value = f"=SUM({col}5:{col}{end})"
        body_cell(c, fmt="0.00" if col in "GJ" else None)
        c.font = F_BOLD
    ws.conditional_formatting.add(f"H5:H{end}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                               end_type="num", end_value=1, end_color="5B9BD5"))
    ws.conditional_formatting.add(f"K5:K{end}", ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                               end_type="num", end_value=0.5, end_color="F4B183"))
    ws.conditional_formatting.add(f"P5:P{end}", CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=FILL_FULL,
                                                           font=Font(name=FONT_NAME, size=10, bold=True, color="C00000")))
    ws.conditional_formatting.add(f"Q5:Q{end}", DataBarRule(start_type="num", start_value=0, end_type="max", color="ED7D31"))
    outline(ws, 4, 1, t, len(cols))
    ws.freeze_panes = "D5"
    ws.cell(row=t + 2, column=1, value="終日続いた貸出不可台数＝この日の貸出不可台数の最小値．2台以上なら赤．予約（最大30分）では説明しにくく，長く動いていない車両がある可能性．").font = F_NOTE
    ws.cell(row=t + 3, column=1, value="推定貸出・返却回数＝15分ごとの「枠内の台数」の減少・増加の合計．15分の間に貸出と返却が相殺された分は数えられないため，実際より少なめ（下限）．").font = F_NOTE


def sheet_hourly(wb, a, w):
    ws = wb.create_sheet("時間帯別")
    sheet_title(ws, "時間帯別の集計", "推定貸出・返却回数は「15分ごとのデータ」シートから数式で計算しています．")
    hr = a["hr"].copy()
    cols = list(hr.columns) + (["気温(℃)", "降水量(mm)", "天気"] if w is not None else [])
    for j, t in enumerate(cols, 1):
        head_cell(ws.cell(row=4, column=j, value=t))
        ws.column_dimensions[get_column_letter(j)].width = 11 if j > 1 else 6
    ws.row_dimensions[4].height = 42
    R = f"'{RAW_SHEET}'!"
    wmap = w.set_index("時") if w is not None else None
    for i, r in enumerate(hr.itertuples(index=False), 5):
        vals = list(r)
        vals[6] = f'=-SUMIFS({R}$L:$L,{R}$C:$C,A{i},{R}$L:$L,"<0")'
        vals[7] = f'=SUMIFS({R}$L:$L,{R}$C:$C,A{i},{R}$L:$L,">0")'
        if wmap is not None:
            vals += [wmap.loc[r.時, "気温(℃)"], wmap.loc[r.時, "降水量(mm)"], wmap.loc[r.時, "天気"]]
        for j, v in enumerate(vals, 1):
            v = None if (not isinstance(v, str) and pd.isna(v)) else v
            align = CENTER if j == 1 or (w is not None and j == len(cols)) else RIGHT
            body_cell(ws.cell(row=i, column=j, value=v), align, "0.0" if j in (3, 4, 5, 6, 9, 10) else None)
    end = 28
    t = end + 1
    body_cell(ws.cell(row=t, column=1, value="計"), CENTER)
    for col in "GH":
        c = ws[f"{col}{t}"]
        c.value = f"=SUM({col}5:{col}{end})"
        body_cell(c)
        c.font = F_BOLD
    ws.conditional_formatting.add(f"C5:C{end}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))
    ws.conditional_formatting.add(f"G5:G{end}", DataBarRule(start_type="num", start_value=0, end_type="max", color="ED7D31"))
    ws.conditional_formatting.add(f"H5:H{end}", DataBarRule(start_type="num", start_value=0, end_type="max", color="4472C4"))
    if w is not None:
        ws.conditional_formatting.add(f"J5:J{end}", CellIsRule(operator="greaterThanOrEqual", formula=["0.5"], fill=FILL_EMPTY))
    outline(ws, 4, 1, t, len(cols))

    # 時間帯 × ステーション（平均の貸出可能台数）
    r0 = t + 3
    h2(ws, r0, 1, "時間帯×ステーション：平均の貸出可能台数（0台＝青，多いほど緑）")
    info = a["info"]
    heat = a["heat"]
    head_cell(ws.cell(row=r0 + 1, column=1, value="時"))
    for j, rr in enumerate(info.itertuples(), 2):
        c = ws.cell(row=r0 + 1, column=j, value=f"{rr.区}\n{rr.ステーション名}\n（容量{int(rr.容量)}）")
        head_cell(c)
        ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, 11)
    ws.row_dimensions[r0 + 1].height = 90
    for i, h in enumerate(range(24), r0 + 2):
        body_cell(ws.cell(row=i, column=1, value=h), CENTER)
        for j, sid in enumerate(info["ステーションID"], 2):
            v = heat.loc[h, sid]
            body_cell(ws.cell(row=i, column=j, value=None if pd.isna(v) else float(v)), CENTER, "0.0")
    last = get_column_letter(len(info) + 1)
    rng = f"B{r0 + 2}:{last}{r0 + 25}"
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="9DC3E6",
                                                      mid_type="num", mid_value=1, mid_color="FFFFFF",
                                                      end_type="max", end_color="63BE7B"))
    outline(ws, r0 + 1, 1, r0 + 25, len(info) + 1)
    ws.freeze_panes = "B5"


def sheet_events(wb, a):
    ws = wb.create_sheet("業者作業の候補")
    sheet_title(ws, "業者作業の候補（再配置・回収・整備の可能性がある変化）",
                f"判定：15分で{BIG_CHANGE}台以上の増減＝大きな増加・減少，台数そのままで貸出不可→貸出可能＝整備らしい動き．利用者の動きと区別できない場合もあります．")
    ev = a["events"]
    widths = [8, 6, 12, 44, 14, 10, 12, 12, 90]
    for j, (t, wd) in enumerate(zip(ev.columns, widths), 1):
        head_cell(ws.cell(row=4, column=j, value=t))
        ws.column_dimensions[get_column_letter(j)].width = wd
    if ev.empty:
        ws.cell(row=5, column=1, value="この日は該当する変化はありませんでした．").font = F_BODY
        return
    fill_kind = {"大きな増加": PatternFill("solid", fgColor="E2EFDA"), "大きな減少": FILL_FULL,
                 "整備らしい動き": PatternFill("solid", fgColor="FFF2CC")}
    for i, r in enumerate(ev.itertuples(index=False), 5):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            body_cell(c, WRAP if j == 9 else LEFT if j in (1, 2, 3, 4, 5) else RIGHT, "+0;-0;0" if j in (6, 7, 8) else None)
        ws.cell(row=i, column=5).fill = fill_kind[r.種類]
    end = 4 + len(ev)
    outline(ws, 4, 1, end, len(ev.columns))
    ws.auto_filter.ref = f"A4:I{end}"
    ws.freeze_panes = "A5"


def sheet_station_graphs(wb, a):
    """25箇所それぞれの15分ごとの推移（make_excel.py と同じ形のグラフ）．"""
    df = a["df"].rename(columns={"station_id": "ステーションID", "num_bikes_available": "貸出可能台数",
                                 "unavailable": "貸出不可台数", "status": "状態"})
    df["日時"] = df["t"].dt.strftime("%Y/%m/%d %H:%M")
    t15, _ = build_series(df, a["info"], "15min")
    t15.index = [s[-5:] for s in t15.index]  # 1日分なので「時:分」だけにする
    dws, col_of = write_chart_data(wb, "グラフ用データ(ステーション)", t15, a["info"])
    sheet_charts(wb, "ステーション別グラフ", dws, col_of, len(t15), a["info"],
                 "ステーション別の15分ごとの推移",
                 "緑＝貸出可能台数，赤＝貸出不可台数，灰色の点線＝容量．緑が0なら空車，緑＋赤が点線に届けば満車．",
                 "時刻", skip=8)


def sheet_help(wb):
    ws = wb.create_sheet("説明")
    sheet_title(ws, "用語と判定の基準")
    rows = [
        ("項目", "意味"),
        ("集計期間", "その日の 0:00 から 23:45 までの15分ごとの記録（96回）．翌日 0:00 の記録は翌日に含める．"),
        ("容量", "ステーションのラック数（物理的に止められる最大台数）"),
        ("貸出可能台数", "今すぐ借りられる自転車の台数"),
        ("空き枠", "今すぐ返却できる枠の数．0なら満車"),
        ("貸出不可台数", "容量 −（貸出可能台数＋空き枠）．故障・電池切れ・予約中（最大30分）などで貸せない自転車の数の目安"),
        ("枠内の台数", "容量 − 空き枠．貸せる自転車と貸せない自転車の合計"),
        ("空車・満車", "空車＝貸出可能台数が0．満車＝空き枠が0．空車時間・満車時間は該当した回数×15分"),
        ("推定貸出・返却回数", "枠内の台数が前の回より減った分を貸出，増えた分を返却として合計したもの．15分の間に貸出と返却が相殺された分は数えられないため実際より少ない（下限）．業者による搬入・持ち出しも含む"),
        ("回転率", "推定貸出回数 ÷ 容量．1日に1台分の枠あたり何回貸し出されたか"),
        ("終日続いた貸出不可台数", "その日の貸出不可台数の最小値．2台以上なら，1日中貸せない自転車が残っていたことになる"),
        ("業者作業の候補", f"大きな増加・減少＝15分で枠内の台数が{BIG_CHANGE}台以上増減．整備らしい動き＝枠内の台数は同じまま貸出不可が減り貸出可能が増えた（バッテリー交換・車両の入れ替え・予約の取り消しのいずれか）"),
        ("文章（ポイント・わかること）", "データから自動で作成．推定や可能性を含むため，卒論に使う際は数値を確認してから使うこと"),
        ("自転車のデータ", "公共交通オープンデータ協議会（ODPT）が公開する HELLO CYCLING の GBFS（station_status / station_information）を15分ごとに取得"),
        ("天気のデータ", "Open-Meteo.com（CC BY 4.0）の1時間ごとの気温・降水量・天気（北緯35.448°，東経139.630°）．観測値ではなく気象モデルの値"),
    ]
    for i, (k, v) in enumerate(rows, 3):
        a_, b_ = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=v)
        for c in (a_, b_):
            c.font, c.border, c.alignment = F_BODY, BORDER, Alignment(vertical="center", wrap_text=True)
        if i == 3:
            head_cell(a_)
            head_cell(b_)
        else:
            a_.fill = FILL_KEY
        ws.row_dimensions[i].height = max(18, 15 * (-(-len(v) // 52)) + 4)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100
    outline(ws, 3, 1, 2 + len(rows), 2)


# ---------------------------------------------------------------------------
def report_path(d):
    e = d + timedelta(days=1)
    return f"{OUT_DIR}/{d.year}/日次レポート_{d.month:02d}.{d.day:02d}-{e.month:02d}.{e.day:02d}.xlsx"


def make_report(d, weather=None):
    a = analyze(d)
    if a is None:
        print(f"  {d}: データがないため作成しません")
        return None
    pa = analyze(d - timedelta(days=1))
    prev_tot = totals(pa) if pa is not None and pa["n_fetch"] >= 80 else None
    w = weather if weather is not None else fetch_weather(d)
    tx = texts(a, w, prev_tot)

    wb = Workbook()
    sheet_overview(wb, a, w, tx, totals(a), prev_tot)
    gd = sheet_graph_data(wb, a, w)
    sheet_graphs(wb, a, w, tx, gd)
    sheet_stations(wb, a)
    sheet_hourly(wb, a, w)
    sheet_events(wb, a)
    sheet_station_graphs(wb, a)
    sheet_raw(wb, a)
    sheet_help(wb)
    # 元データのシートは最後に回す
    data_sheets = [s for s in wb._sheets if s.title.startswith("グラフ用データ")]
    wb._sheets = [s for s in wb._sheets if s not in data_sheets] + data_sheets
    wb.active = 0
    # 印刷するときは横向き・横幅を1ページに収める
    for s in wb.worksheets:
        s.page_setup.orientation = "landscape"
        s.page_setup.paperSize = s.PAPERSIZE_A4
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0

    path = report_path(d)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    print(f"✅ {path}（取得 {a['n_fetch']} 回，天気 {'あり' if w is not None else 'なし'}）")
    return path


def main():
    today = datetime.now(JST).date()
    if len(sys.argv) > 1:
        for s in sys.argv[1:]:
            make_report(datetime.strptime(s, "%Y%m%d").date())
        return
    # まだレポートがない「終わった日」をすべて作る
    days = sorted({datetime.strptime(os.path.basename(p)[-12:-4], "%Y%m%d").date()
                   for p in glob.glob(f"{DATA_DIR}/yokohama_stations_15min_log_raw_*.csv")})
    for d in days:
        if d < today and not os.path.exists(report_path(d)):
            make_report(d)


if __name__ == "__main__":
    main()
